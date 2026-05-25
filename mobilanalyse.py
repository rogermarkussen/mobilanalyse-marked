# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "marimo>=0.23.3",
#   "matplotlib>=3.8.0",
#   "openpyxl>=3.1.0",
#   "pyarrow>=16.0.0",
#   "polars>=1.0.0",
# ]
# ///

import marimo

__generated_with = "0.23.8"
app = marimo.App(width="full")


@app.cell
def _():
    from base64 import b64encode
    from io import BytesIO
    from json import dumps
    from pathlib import Path

    import matplotlib.pyplot as plt
    import marimo as mo
    import polars as pl
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    try:
        from pyodide.http import pyfetch
    except ImportError:
        pyfetch = None
    return (
        Alignment,
        Border,
        BytesIO,
        Font,
        Path,
        PatternFill,
        Side,
        Workbook,
        b64encode,
        dumps,
        mo,
        pl,
        plt,
        pyfetch,
    )


@app.cell
async def _(Path, pl, pyfetch):
    if pyfetch is None:
        parquet_path = Path(__file__).parent / "data" / "mobil.parquet"
    else:
        response = await pyfetch("../data/mobil.parquet")
        parquet_path = Path("/tmp/mobil.parquet")
        parquet_path.write_bytes(await response.bytes())

    df = pl.scan_parquet(parquet_path)
    return (df,)


@app.cell
def _(
    Alignment,
    Border,
    BytesIO,
    Font,
    PatternFill,
    Side,
    Workbook,
    b64encode,
    dumps,
    pl,
):
    def export_menu(
        _menu_id,
        _data,
        _excel_filename,
        _title,
        _sheet_name,
        _figure,
        _png_filename,
    ):
        _is_trend_export = "serie" in _data.columns
        if _is_trend_export:
            _index_columns = [
                _column for _column in ["ms", "tilbyder"] if _column in _data.columns
            ]
            _history_years = sorted(
                _data.filter(pl.col("serie") == "Historikk")["ar"].unique().to_list()
            )
            _last_history_year = max(_history_years)
            _trend_years = sorted(
                _data.filter(
                    (pl.col("serie") == "Lineær trend")
                    & (pl.col("ar") > _last_history_year)
                )["ar"].unique().to_list()
            )
            _export_data = (
                _data.filter(
                    (pl.col("serie") == "Historikk")
                    | (
                        (pl.col("serie") == "Lineær trend")
                        & (pl.col("ar") > _last_history_year)
                    )
                )
                .with_columns(
                    (pl.col("markedsandel") / 100).alias("markedsandel"),
                    pl.col("ar").cast(pl.Utf8),
                )
                .pivot(
                    values="markedsandel",
                    index=_index_columns,
                    on="ar",
                    aggregate_function="first",
                )
                .select(
                    _index_columns
                    + [str(_year) for _year in _history_years + _trend_years]
                )
                .sort(_index_columns)
            )
            _years = _history_years + _trend_years
        else:
            _index_columns = [
                _column for _column in ["ms", "tilbyder"] if _column in _data.columns
            ]
            _years = sorted(_data["ar"].unique().to_list())
            _export_data = (
                _data.with_columns(
                    (pl.col("markedsandel") / 100).alias("markedsandel"),
                    pl.col("ar").cast(pl.Utf8),
                )
                .pivot(
                    values="markedsandel",
                    index=_index_columns,
                    on="ar",
                    aggregate_function="first",
                )
                .select(_index_columns + [str(_year) for _year in _years])
                .sort(_index_columns)
            )
        _headers = _export_data.columns
        _header_row = 4 if _is_trend_export else 3
        _first_data_row = _header_row + 1
        _workbook = Workbook()
        _sheet = _workbook.active
        _sheet.title = _sheet_name[:31]

        _sheet["A1"] = _title
        _sheet["A1"].font = Font(bold=True, size=14, color="0B2B66")
        _sheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=max(1, len(_headers)),
        )

        _sheet.append([])
        if _is_trend_export:
            _sheet.append([""] * len(_headers))
            _history_start = len(_index_columns) + 1
            _history_end = _history_start + len(_history_years) - 1
            _trend_start = _history_end + 1
            _trend_end = _trend_start + len(_trend_years) - 1
            if _history_years:
                _sheet.merge_cells(
                    start_row=3,
                    start_column=_history_start,
                    end_row=3,
                    end_column=_history_end,
                )
                _sheet.cell(row=3, column=_history_start).value = "Historikk"
            if _trend_years:
                _sheet.merge_cells(
                    start_row=3,
                    start_column=_trend_start,
                    end_row=3,
                    end_column=_trend_end,
                )
                _sheet.cell(row=3, column=_trend_start).value = "Lineær trend"
        _sheet.append(_headers)
        for _row in _export_data.iter_rows():
            _sheet.append(list(_row))

        _header_fill = PatternFill("solid", fgColor="0B2B66")
        _header_font = Font(bold=True, color="FFFFFF")
        _thin = Side(style="thin", color="D9D9D9")
        _medium_blue = Side(style="medium", color="0B2B66")
        _border = Border(bottom=_thin)

        if _is_trend_export:
            _history_fill = PatternFill("solid", fgColor="E9F2DF")
            _trend_fill = PatternFill("solid", fgColor="DDEBF7")
            for _cell in _sheet[3]:
                _cell.fill = _history_fill
                _cell.font = Font(bold=True, color="000000")
                _cell.alignment = Alignment(horizontal="center")
                _cell.border = Border(top=_medium_blue, bottom=_thin)
            for _column_index in range(_trend_start, _trend_end + 1):
                _cell = _sheet.cell(row=3, column=_column_index)
                _cell.fill = _trend_fill
                _cell.border = Border(top=_medium_blue, bottom=_thin)
            for _row_index in range(3, _sheet.max_row + 1):
                _left_cell = _sheet.cell(row=_row_index, column=_trend_start)
                _left_cell.border = Border(
                    left=_medium_blue,
                    right=_left_cell.border.right,
                    top=_left_cell.border.top,
                    bottom=_left_cell.border.bottom,
                )

        for _cell in _sheet[_header_row]:
            _cell.fill = _header_fill
            _cell.font = _header_font
            _cell.alignment = Alignment(horizontal="center")
            if _is_trend_export and _cell.column >= _trend_start:
                _cell.fill = PatternFill("solid", fgColor="1F4E79")

        for _row in _sheet.iter_rows(
            min_row=_first_data_row,
            max_row=_sheet.max_row,
            max_col=_sheet.max_column,
        ):
            for _cell in _row:
                _cell.border = _border
                if _cell.column_letter == "A":
                    _cell.alignment = Alignment(horizontal="left")
                if _headers[_cell.column - 1] in [str(_year) for _year in _years]:
                    _cell.number_format = "0.0%"

        _table_ref = f"A{_header_row}:{_sheet.cell(_sheet.max_row, _sheet.max_column).coordinate}"
        for _row_index in range(_first_data_row, _sheet.max_row + 1):
            if (_row_index - _first_data_row) % 2 == 0:
                for _cell in _sheet[_row_index]:
                    _cell.fill = PatternFill("solid", fgColor="EEF3F8")
            if _is_trend_export:
                for _column_index in range(len(_index_columns) + 1, _history_end + 1):
                    _sheet.cell(row=_row_index, column=_column_index).fill = PatternFill(
                        "solid",
                        fgColor="F3F8EC",
                    )
                for _column_index in range(_trend_start, _trend_end + 1):
                    _sheet.cell(row=_row_index, column=_column_index).fill = PatternFill(
                        "solid",
                        fgColor="EAF3FB",
                    )
                _left_cell = _sheet.cell(row=_row_index, column=_trend_start)
                _left_cell.border = Border(
                    left=_medium_blue,
                    right=_left_cell.border.right,
                    top=_left_cell.border.top,
                    bottom=_left_cell.border.bottom,
                )
        _sheet.freeze_panes = f"A{_first_data_row}"
        _sheet.auto_filter.ref = _table_ref

        for _column_index, _header in enumerate(_headers, start=1):
            _max_length = len(str(_header))
            for _row_index in range(_first_data_row, _sheet.max_row + 1):
                _value = _sheet.cell(row=_row_index, column=_column_index).value
                _max_length = max(
                    _max_length,
                    len(str(_value)) if _value is not None else 0,
            )
            _sheet.column_dimensions[
                _sheet.cell(row=_header_row, column=_column_index).column_letter
            ].width = min(
                max(_max_length + 2, 12),
                28,
            )

        _buffer = BytesIO()
        _workbook.save(_buffer)
        _excel_payload = b64encode(_buffer.getvalue()).decode("ascii")
        _png_buffer = BytesIO()
        _figure.savefig(
            _png_buffer,
            format="png",
            dpi=180,
            bbox_inches="tight",
            facecolor="white",
        )
        _png_payload = b64encode(_png_buffer.getvalue()).decode("ascii")

        _menu_html = f"""
            <details id="{_menu_id}" class="export-menu">
                <summary aria-label="Eksporter figur" title="Eksporter figur">
                    <svg viewBox="0 0 24 24" aria-hidden="true">
                        <path d="M4 7h16"></path>
                        <path d="M4 12h16"></path>
                        <path d="M4 17h16"></path>
                    </svg>
                </summary>
                <div class="export-menu-items">
                    <a
                        href="data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,{_excel_payload}"
                        download="{_excel_filename}"
                    >Excel</a>
                    <a
                        href="data:image/png;base64,{_png_payload}"
                        download="{_png_filename}"
                    >PNG</a>
                </div>
            </details>
            """

        return (
            "<script>"
            f"const menu = document.getElementById({dumps(_menu_id)});"
            f"if (menu) menu.outerHTML = {dumps(_menu_html)};"
            "</script>"
        )

    def export_menu_placeholder(_menu_id):
        return f"""
            <details id="{_menu_id}" class="export-menu export-menu-loading">
                <summary aria-label="Eksport klargjøres" title="Eksport klargjøres">
                    <svg viewBox="0 0 24 24" aria-hidden="true">
                        <path d="M4 7h16"></path>
                        <path d="M4 12h16"></path>
                        <path d="M4 17h16"></path>
                    </svg>
                </summary>
                <div class="export-menu-items">
                    <span>Klargjør eksport...</span>
                </div>
            </details>
            """

    return export_menu, export_menu_placeholder


@app.cell
def _(mo):
    mo.Html(
        """
        <style>
            .figure-heading-title {
                align-items: center;
                display: inline-flex;
                font-size: 1.25rem;
                font-weight: 700;
                line-height: 30px;
            }

            .export-menu {
                display: inline-flex;
                position: relative;
            }

            .figure-export-wrap {
                display: inline-block;
                position: relative;
            }

            .figure-export-image {
                display: block;
                height: auto;
                max-width: 100%;
            }

            .figure-export-wrap > .export-menu {
                position: absolute;
                right: 13%;
                top: 6%;
                z-index: 10;
            }

            .export-menu summary {
                align-items: center !important;
                border-radius: 4px !important;
                color: #5f6f82 !important;
                cursor: pointer;
                display: inline-flex !important;
                height: 42px !important;
                justify-content: center !important;
                margin: 0 !important;
                list-style: none;
                width: 42px !important;
            }

            .export-menu summary::-webkit-details-marker {
                display: none;
            }

            .export-menu summary:hover {
                background: #f4f7fa !important;
                color: #0b2b66 !important;
            }

            .export-menu svg {
                color: currentColor !important;
                fill: none !important;
                height: 28px !important;
                stroke: currentColor !important;
                stroke-linecap: round !important;
                stroke-linejoin: round !important;
                stroke-width: 1.8 !important;
                width: 28px !important;
            }

            .export-menu-items {
                background: #ffffff;
                border: 1px solid #d8e0eb;
                border-radius: 6px;
                box-shadow: 0 8px 24px rgba(11, 43, 102, 0.14);
                display: grid;
                min-width: 96px;
                padding: 4px;
                position: absolute;
                right: 0;
                top: 28px;
                z-index: 10;
            }

            .export-menu-items a {
                border-radius: 4px;
                color: #1f2937;
                font-size: 0.86rem;
                line-height: 1.2;
                padding: 7px 10px;
                text-decoration: none;
            }

            .export-menu-loading summary {
                opacity: 0.45;
            }

            .export-menu-items span {
                color: #64748b;
                font-size: 0.82rem;
                padding: 7px 10px;
                white-space: nowrap;
            }

            .export-menu-items a:hover {
                background: #f4f7fa;
                color: #0b2b66;
            }
        </style>
        <div style="margin-bottom: 20px;">
            <div style="font-size: 2.2rem; font-weight: 700; color: #0b2b66;">
                Mobilanalyse marked
            </div>
        </div>
        """
    )
    return


@app.cell
def _(
    export_menu,
    fig1_abonnement_fig,
    fig1_omsetning_fig,
    fig2_abonnement_fig,
    fig2_omsetning_fig,
    fig3_business_fig,
    fig3_private_fig,
    market_share_abonnement,
    market_share_abonnement_projection,
    market_share_abonnement_segment,
    market_share_omsetning,
    market_share_omsetning_projection,
    mo,
    pl,
    projection_export_data,
):
    _updates = [
        export_menu(
            "fig1-abonnement-export",
            market_share_abonnement,
            "figur-1-abonnement.xlsx",
            "Figur 1 - Markedsandeler basert på abonnement",
            "Abonnement",
            fig1_abonnement_fig,
            "figur-1-abonnement.png",
        ),
        export_menu(
            "fig1-omsetning-export",
            market_share_omsetning,
            "figur-1-omsetning.xlsx",
            "Figur 1 - Markedsandeler basert på omsetning",
            "Omsetning",
            fig1_omsetning_fig,
            "figur-1-omsetning.png",
        ),
        export_menu(
            "fig2-abonnement-export",
            projection_export_data(
                market_share_abonnement,
                market_share_abonnement_projection,
            ),
            "figur-2-abonnement-trend.xlsx",
            "Figur 2 - Lineær trend basert på abonnement",
            "Abonnement trend",
            fig2_abonnement_fig,
            "figur-2-abonnement-trend.png",
        ),
        export_menu(
            "fig2-omsetning-export",
            projection_export_data(
                market_share_omsetning,
                market_share_omsetning_projection,
            ),
            "figur-2-omsetning-trend.xlsx",
            "Figur 2 - Lineær trend basert på omsetning",
            "Omsetning trend",
            fig2_omsetning_fig,
            "figur-2-omsetning-trend.png",
        ),
        export_menu(
            "fig3-privat-export",
            market_share_abonnement_segment.filter(pl.col("ms") == "Privat"),
            "figur-3-privat.xlsx",
            "Figur 3 - Abonnement i privatmarkedet",
            "Privat",
            fig3_private_fig,
            "figur-3-privat.png",
        ),
        export_menu(
            "fig3-bedrift-export",
            market_share_abonnement_segment.filter(pl.col("ms") == "Bedrift"),
            "figur-3-bedrift.xlsx",
            "Figur 3 - Abonnement i bedriftsmarkedet",
            "Bedrift",
            fig3_business_fig,
            "figur-3-bedrift.png",
        ),
    ]
    mo.Html("\n".join(_updates))
    return


@app.cell
def _(BytesIO, b64encode, mo):
    def figure_with_export(_figure, _menu):
        _buffer = BytesIO()
        _figure.savefig(
            _buffer,
            format="png",
            dpi=120,
            bbox_inches="tight",
            facecolor="white",
        )
        _payload = b64encode(_buffer.getvalue()).decode("ascii")
        return mo.Html(
            f"""
            <div class="figure-export-wrap">
                {_menu}
                <img class="figure-export-image" src="data:image/png;base64,{_payload}" alt="">
            </div>
            """
        )

    return (figure_with_export,)


@app.cell
def _(mo):
    mo.md("""
    ## 1 - Utvikling i markedsandeler
    """)
    return


@app.cell
def _(df, pl):
    _provider_name = pl.col("fusnavn").str.to_lowercase()
    _provider_group = (
        pl.when(_provider_name.str.contains("telenor"))
        .then(pl.lit("Telenor"))
        .when(_provider_name.str.contains("telia"))
        .then(pl.lit("Telia"))
        .when(
            _provider_name.is_in(["ice communication norge", "lyse tele"])
            | _provider_name.str.contains("lyse")
            | _provider_name.str.contains("ice")
        )
        .then(pl.lit("Lyse Tele (Ice)"))
        .otherwise(pl.lit("Øvrige"))
    )

    market_share_abonnement = (
        df.filter(
            (pl.col("dk") == "Mobiltelefoni")
            & (pl.col("hg") == "Abonnement")
            & (pl.col("n1").is_in(["Fakturert", "Kontantkort"]))
            & (pl.col("n2") == "Ingen")
            & (pl.col("tp") == "Sum")
            & (pl.col("sk") == "Sluttbruker")
            & (pl.col("delar") == "Helår")
        )
        .with_columns(tilbyder=_provider_group)
        .group_by("ar", "tilbyder")
        .agg(pl.col("svar").sum().alias("abonnement"))
        .with_columns(
            markedsandel=pl.col("abonnement")
            * 100
            / pl.col("abonnement").sum().over("ar")
        )
        .select("ar", "tilbyder", "markedsandel")
        .sort("ar", "tilbyder")
        .collect()
    )
    return (market_share_abonnement,)


@app.cell
def _(df, pl):
    _provider_name = pl.col("fusnavn").str.to_lowercase()
    _provider_group = (
        pl.when(_provider_name.str.contains("telenor"))
        .then(pl.lit("Telenor"))
        .when(_provider_name.str.contains("telia"))
        .then(pl.lit("Telia"))
        .when(
            _provider_name.is_in(["ice communication norge", "lyse tele"])
            | _provider_name.str.contains("lyse")
            | _provider_name.str.contains("ice")
        )
        .then(pl.lit("Lyse Tele (Ice)"))
        .otherwise(pl.lit("Øvrige"))
    )

    market_share_omsetning = (
        df.filter(
            (pl.col("dk") == "Mobiltelefoni")
            & (pl.col("hg") == "Inntekter")
            & (pl.col("tp") == "Sum")
            & (pl.col("sk") == "Sluttbruker")
            & (pl.col("delar") == "Helår")
        )
        .with_columns(tilbyder=_provider_group)
        .group_by("ar", "tilbyder")
        .agg(pl.col("svar").sum().alias("omsetning"))
        .with_columns(
            markedsandel=pl.col("omsetning")
            * 100
            / pl.col("omsetning").sum().over("ar")
        )
        .select("ar", "tilbyder", "markedsandel")
        .sort("ar", "tilbyder")
        .collect()
    )
    return (market_share_omsetning,)


@app.cell
def _(
    export_menu_placeholder,
    figure_with_export,
    market_share_abonnement,
    market_share_omsetning,
    mo,
    pl,
    plt,
):
    _colors = {
        "Telenor": "#156082",
        "Telia": "#7030a0",
        "Lyse Tele (Ice)": "#ffc000",
        "Øvrige": "#00b050",
    }
    _order = ["Telenor", "Telia", "Lyse Tele (Ice)", "Øvrige"]

    def _percent(_value):
        return f"{_value:.1f} %".replace(".", ",")

    def _plot_market_share(_data, _upper_y):
        _fig, _ax = plt.subplots(figsize=(7.2, 4.8), dpi=120)
        for _provider in _order:
            _provider_data = _data.filter(_data["tilbyder"] == _provider).sort("ar")
            if _provider_data.is_empty():
                continue

            _xs = _provider_data["ar"].to_list()
            _ys = _provider_data["markedsandel"].to_list()
            _ax.plot(
                _xs,
                _ys,
                color=_colors[_provider],
                linewidth=3.0,
                solid_capstyle="round",
                label=_provider,
            )

            if len(_ys) >= 2:
                _ax.annotate(
                    _percent(_ys[-2]),
                    xy=(_xs[-2], _ys[-2]),
                    xytext=(10, 6),
                    textcoords="offset points",
                    color="#404040",
                    fontsize=9,
                )
            _ax.annotate(
                _percent(_ys[-1]),
                xy=(_xs[-1], _ys[-1]),
                xytext=(10, 6),
                textcoords="offset points",
                color="#404040",
                fontsize=9,
            )

        _ax.set_ylim(0, _upper_y)
        _ax.set_yticks(range(0, _upper_y + 1, 10))
        _ax.set_yticklabels([_percent(_value) for _value in range(0, _upper_y + 1, 10)])
        _ax.set_xticks(sorted(_data["ar"].unique().to_list()))
        _ax.grid(axis="y", color="#d9d9d9", linewidth=0.8)
        _ax.grid(axis="x", visible=False)
        _ax.spines[["top", "right", "left"]].set_visible(False)
        _ax.spines["bottom"].set_color("#d9d9d9")
        _ax.tick_params(axis="both", colors="#555555", length=0)
        _ax.legend(
            loc="lower center",
            bbox_to_anchor=(0.5, -0.18),
            ncol=4,
            frameon=False,
            fontsize=10,
        )
        _fig.tight_layout()
        return _fig

    def _latest_change(_data, _provider):
        _provider_data = _data.filter(pl.col("tilbyder") == _provider).sort("ar")
        _periods = _provider_data["ar"].to_list()
        _values = _provider_data["markedsandel"].to_list()
        if len(_values) < 2:
            return None
        return {
            "previous_period": _periods[-2],
            "current_period": _periods[-1],
            "previous": _values[-2],
            "current": _values[-1],
            "delta": _values[-1] - _values[-2],
        }

    def _direction(_change):
        if _change["delta"] > 0.05:
            return "øker"
        if _change["delta"] < -0.05:
            return "går ned"
        return "er om lag uendret"

    def _pair_text(_provider):
        _ab = _latest_change(market_share_abonnement, _provider)
        _om = _latest_change(market_share_omsetning, _provider)
        if _ab is None or _om is None:
            return ""
        return (
            f"<strong>{_provider}</strong> {_direction(_om)} i omsetningsandel "
            f"(<strong>{_percent(_om['previous'])}→{_percent(_om['current'])}</strong>) "
            f"og {_direction(_ab)} i abonnementsandel "
            f"(<strong>{_percent(_ab['previous'])}→{_percent(_ab['current'])}</strong>)."
        )

    def _last_two_periods():
        _all_years = sorted(
            set(
                market_share_abonnement["ar"].to_list()
                + market_share_omsetning["ar"].to_list()
            )
        )
        if len(_all_years) < 2:
            return None, None
        return _all_years[-2], _all_years[-1]

    _prev_period, _curr_period = _last_two_periods()
    _period_text = (
        f"Endring fra {_prev_period} til {_curr_period}"
        if _prev_period is not None and _curr_period is not None
        else "Endring mellom siste perioder"
    )

    fig1_abonnement_fig = _plot_market_share(market_share_abonnement, 60)
    fig1_omsetning_fig = _plot_market_share(market_share_omsetning, 60)
    _abonnement_export = export_menu_placeholder("fig1-abonnement-export")
    _omsetning_export = export_menu_placeholder("fig1-omsetning-export")

    _summary = mo.Html(
        f"""
        <div style="
            margin: 8px auto 0;
            max-width: 820px;
            padding: 18px 26px;
            border: 1px solid #0b2b66;
            border-radius: 28px;
            background: #f6c7a7;
            color: #000;
            font-size: 1.05rem;
            line-height: 1.45;
            text-align: center;
        ">
            <div style="margin: 0 0 12px; font-size: 1.2rem; font-weight: 700;">{_period_text}</div>
            {_pair_text("Lyse Tele (Ice)")}<br>
            {_pair_text("Telenor")}<br>
            {_pair_text("Telia")}<br>
            {_pair_text("Øvrige")}
        </div>
        """
    )

    mo.vstack(
        [
            mo.hstack(
                [
                    mo.vstack(
                        [
                            mo.Html(
                                '<div class="figure-heading-title">Basert på abonnement</div>'
                            ),
                            figure_with_export(fig1_abonnement_fig, _abonnement_export),
                        ],
                        gap=0.5,
                    ),
                    mo.vstack(
                        [
                            mo.Html(
                                '<div class="figure-heading-title">Basert på omsetning</div>'
                            ),
                            figure_with_export(fig1_omsetning_fig, _omsetning_export),
                        ],
                        gap=0.5,
                    ),
                ],
                justify="center",
                gap=2,
            ),
            _summary,
        ],
        gap=2,
    )
    return fig1_abonnement_fig, fig1_omsetning_fig


@app.cell
def _(mo):
    mo.Html(
        """
        <div style="
            width: 100%;
            margin: 46px 0 24px;
            border-top: 3px solid #0b2b66;
        "></div>
        """
    )
    return


@app.cell
def _(mo):
    mo.md("""
    ## 2 - Lineær trend i markedsandeler
    """)
    return


@app.cell
def _(market_share_abonnement, market_share_omsetning, pl):
    def _linear_projection(_data, _periods_ahead=3):
        _rows = []
        for _provider in ["Telenor", "Telia", "Lyse Tele (Ice)", "Øvrige"]:
            _provider_data = _data.filter(pl.col("tilbyder") == _provider).sort("ar")
            if _provider_data.height < 2:
                continue

            _xs = [float(_x) for _x in _provider_data["ar"].to_list()]
            _ys = [float(_y) for _y in _provider_data["markedsandel"].to_list()]
            _x_mean = sum(_xs) / len(_xs)
            _y_mean = sum(_ys) / len(_ys)
            _denominator = sum((_x - _x_mean) ** 2 for _x in _xs)
            _slope = (
                sum((_x - _x_mean) * (_y - _y_mean) for _x, _y in zip(_xs, _ys))
                / _denominator
                if _denominator
                else 0
            )
            _last_year = int(max(_xs))
            _last_value = _ys[-1]

            for _year in range(_last_year, _last_year + _periods_ahead + 1):
                _rows.append(
                    {
                        "ar": _year,
                        "tilbyder": _provider,
                        "markedsandel": _last_value + _slope * (_year - _last_year),
                    }
                )

        return pl.DataFrame(_rows)

    market_share_abonnement_projection = _linear_projection(market_share_abonnement)
    market_share_omsetning_projection = _linear_projection(market_share_omsetning)
    return (
        market_share_abonnement_projection,
        market_share_omsetning_projection,
    )


@app.cell
def _(
    export_menu_placeholder,
    figure_with_export,
    market_share_abonnement,
    market_share_abonnement_projection,
    market_share_omsetning,
    market_share_omsetning_projection,
    mo,
    pl,
    plt,
):
    _colors = {
        "Telenor": "#156082",
        "Telia": "#7030a0",
        "Lyse Tele (Ice)": "#ffc000",
        "Øvrige": "#00b050",
    }
    _order = ["Telenor", "Telia", "Lyse Tele (Ice)", "Øvrige"]

    def _percent(_value):
        return f"{_value:.1f} %".replace(".", ",")

    def _plot_projection(_actual, _projection, _title, _upper_y):
        _fig, _ax = plt.subplots(figsize=(7.8, 4.8), dpi=120)
        _last_actual_year = int(_actual["ar"].max())

        for _provider in _order:
            _actual_provider = _actual.filter(pl.col("tilbyder") == _provider).sort("ar")
            _projection_provider = _projection.filter(pl.col("tilbyder") == _provider).sort("ar")
            if _actual_provider.is_empty() or _projection_provider.is_empty():
                continue

            _actual_xs = _actual_provider["ar"].to_list()
            _actual_ys = _actual_provider["markedsandel"].to_list()
            _projection_xs = _projection_provider["ar"].to_list()
            _projection_ys = _projection_provider["markedsandel"].to_list()

            _ax.plot(
                _actual_xs,
                _actual_ys,
                color=_colors[_provider],
                linewidth=2.5,
                solid_capstyle="round",
                label=_provider,
            )
            _ax.plot(
                _projection_xs,
                _projection_ys,
                color=_colors[_provider],
                linewidth=1.8,
                linestyle=":",
                label="_nolegend_",
            )

            _latest_actual = _actual_provider.filter(pl.col("ar") == _last_actual_year)
            if not _latest_actual.is_empty():
                _ax.annotate(
                    _percent(_latest_actual["markedsandel"][0]),
                    xy=(_last_actual_year, _latest_actual["markedsandel"][0]),
                    xytext=(8, 7),
                    textcoords="offset points",
                    color="#404040",
                    fontsize=8,
                )

            _ax.annotate(
                _percent(_projection_ys[-1]),
                xy=(_projection_xs[-1], _projection_ys[-1]),
                xytext=(8, 0),
                textcoords="offset points",
                color="#202020",
                fontsize=8,
                bbox={
                    "boxstyle": "square,pad=0.35",
                    "facecolor": "white",
                    "edgecolor": _colors[_provider],
                    "linewidth": 0.8,
                },
            )

        _ax.set_ylim(0, _upper_y)
        _ax.set_yticks(range(0, _upper_y + 1, 10))
        _ax.set_yticklabels([_percent(_value) for _value in range(0, _upper_y + 1, 10)])
        _ax.set_xticks(
            sorted(set(_actual["ar"].unique().to_list() + _projection["ar"].unique().to_list()))
        )
        _ax.grid(axis="y", color="#d9d9d9", linewidth=0.8)
        _ax.grid(axis="x", visible=False)
        _ax.spines[["top", "right", "left"]].set_visible(False)
        _ax.spines["bottom"].set_color("#d9d9d9")
        _ax.tick_params(axis="both", colors="#555555", length=0)
        _ax.legend(
            loc="lower center",
            bbox_to_anchor=(0.5, -0.26),
            ncol=4,
            frameon=False,
            fontsize=8,
        )
        _fig.tight_layout()
        return _fig

    def _threshold_text(_data, _provider, _threshold):
        _provider_data = _data.filter(pl.col("tilbyder") == _provider).sort("ar")
        _xs = [float(_x) for _x in _provider_data["ar"].to_list()]
        _ys = [float(_y) for _y in _provider_data["markedsandel"].to_list()]
        _x_mean = sum(_xs) / len(_xs)
        _y_mean = sum(_ys) / len(_ys)
        _denominator = sum((_x - _x_mean) ** 2 for _x in _xs)
        _slope = (
            sum((_x - _x_mean) * (_y - _y_mean) for _x, _y in zip(_xs, _ys))
            / _denominator
            if _denominator
            else 0
        )
        if _slope <= 0:
            return f"Lyse Tele (Ice) når ikke {_threshold:.0f} % omsetningsandel med lineær trend."
        _last_year = max(_xs)
        _last_value = _ys[-1]
        _year = _last_year + ((_threshold - _last_value) / _slope)
        return f"Lyse Tele (Ice) når {_threshold:.0f} % av omsetningen rundt {int(_year + 0.999)}."

    def projection_export_data(_actual, _projection):
        _actual_export = _actual.with_columns(
            pl.lit("Historikk").alias("serie"),
            pl.col("ar").cast(pl.Int64),
        )
        _projection_export = _projection.with_columns(
            pl.lit("Lineær trend").alias("serie"),
            pl.col("ar").cast(pl.Int64),
        )
        return (
            pl.concat([_actual_export, _projection_export])
            .select("serie", "ar", "tilbyder", "markedsandel")
            .sort("serie", "ar", "tilbyder")
        )

    fig2_abonnement_fig = _plot_projection(
        market_share_abonnement,
        market_share_abonnement_projection,
        "Abonnement",
        60,
    )
    fig2_omsetning_fig = _plot_projection(
        market_share_omsetning,
        market_share_omsetning_projection,
        "Omsetning",
        60,
    )
    _abonnement_projection_export = export_menu_placeholder("fig2-abonnement-export")
    _omsetning_projection_export = export_menu_placeholder("fig2-omsetning-export")
    _forecast_note = mo.Html(
        f"""
        <div style="
            margin: 18px auto 0;
            width: fit-content;
            max-width: 520px;
            padding: 22px 44px;
            border: 1px solid #0b2b66;
            border-radius: 999px;
            background: #c9570d;
            color: white;
            font-size: 1.1rem;
            font-weight: 700;
            line-height: 1.35;
            text-align: center;
        ">
            {_threshold_text(market_share_omsetning, "Lyse Tele (Ice)", 20)}
        </div>
        """
    )

    mo.vstack(
        [
            mo.hstack(
                [
                    mo.vstack(
                        [
                            mo.Html(
                                '<div class="figure-heading-title">Abonnement</div>'
                            ),
                            figure_with_export(fig2_abonnement_fig, _abonnement_projection_export),
                        ],
                        gap=0.5,
                    ),
                    mo.vstack(
                        [
                            mo.Html(
                                '<div class="figure-heading-title">Omsetning</div>'
                            ),
                            figure_with_export(fig2_omsetning_fig, _omsetning_projection_export),
                        ],
                        gap=0.5,
                    ),
                ],
                justify="center",
                gap=2,
            ),
            _forecast_note,
        ],
        gap=1,
    )
    return fig2_abonnement_fig, fig2_omsetning_fig, projection_export_data


@app.cell
def _(mo):
    mo.Html(
        """
        <div style="
            width: 100%;
            margin: 46px 0 24px;
            border-top: 3px solid #0b2b66;
        "></div>
        """
    )
    return


@app.cell
def _(mo):
    mo.md("""
    ## 3 - Abonnement fordelt på privat og bedrift
    """)
    return


@app.cell
def _(df, pl):
    _provider_name = pl.col("fusnavn").str.to_lowercase()
    _provider_group = (
        pl.when(_provider_name.str.contains("telenor"))
        .then(pl.lit("Telenor"))
        .when(_provider_name.str.contains("telia"))
        .then(pl.lit("Telia"))
        .when(
            _provider_name.is_in(["ice communication norge", "lyse tele"])
            | _provider_name.str.contains("lyse")
            | _provider_name.str.contains("ice")
        )
        .then(pl.lit("Lyse Tele (Ice)"))
        .otherwise(pl.lit("Øvrige"))
    )

    market_share_abonnement_segment = (
        df.filter(
            (pl.col("dk") == "Mobiltelefoni")
            & (pl.col("hg") == "Abonnement")
            & (pl.col("ms").is_in(["Privat", "Bedrift"]))
            & (pl.col("n1").is_in(["Fakturert", "Kontantkort"]))
            & (pl.col("n2") == "Ingen")
            & (pl.col("tp") == "Sum")
            & (pl.col("sk") == "Sluttbruker")
            & (pl.col("delar") == "Helår")
        )
        .with_columns(tilbyder=_provider_group)
        .group_by("ar", "ms", "tilbyder")
        .agg(pl.col("svar").sum().alias("abonnement"))
        .with_columns(
            markedsandel=pl.col("abonnement")
            * 100
            / pl.col("abonnement").sum().over(["ar", "ms"])
        )
        .select("ar", "ms", "tilbyder", "markedsandel")
        .sort("ms", "ar", "tilbyder")
        .collect()
    )
    return (market_share_abonnement_segment,)


@app.cell
def _(
    export_menu_placeholder,
    figure_with_export,
    market_share_abonnement_segment,
    mo,
    pl,
    plt,
):
    _colors = {
        "Telenor": "#156082",
        "Telia": "#7030a0",
        "Lyse Tele (Ice)": "#ffc000",
        "Øvrige": "#00b050",
    }
    _order = ["Telenor", "Telia", "Lyse Tele (Ice)", "Øvrige"]

    def _percent(_value):
        return f"{_value:.1f} %".replace(".", ",")

    def _plot_segment(_segment, _upper_y=60):
        _data = market_share_abonnement_segment.filter(pl.col("ms") == _segment)
        _fig, _ax = plt.subplots(figsize=(7.2, 4.8), dpi=120)

        for _provider in _order:
            _provider_data = _data.filter(pl.col("tilbyder") == _provider).sort("ar")
            if _provider_data.is_empty():
                continue

            _xs = _provider_data["ar"].to_list()
            _ys = _provider_data["markedsandel"].to_list()
            _ax.plot(
                _xs,
                _ys,
                color=_colors[_provider],
                linewidth=3.0,
                solid_capstyle="round",
                label=_provider,
            )

            if len(_ys) >= 2:
                _ax.annotate(
                    _percent(_ys[-2]),
                    xy=(_xs[-2], _ys[-2]),
                    xytext=(8, 7),
                    textcoords="offset points",
                    color="#404040",
                    fontsize=9,
                )
            _ax.annotate(
                _percent(_ys[-1]),
                xy=(_xs[-1], _ys[-1]),
                xytext=(10, 0),
                textcoords="offset points",
                color="#404040",
                fontsize=9,
            )

        _ax.set_ylim(0, _upper_y)
        _ax.set_yticks(range(0, _upper_y + 1, 10))
        _ax.set_yticklabels([_percent(_value) for _value in range(0, _upper_y + 1, 10)])
        _ax.set_xticks(sorted(_data["ar"].unique().to_list()))
        _ax.grid(axis="y", color="#d9d9d9", linewidth=0.8)
        _ax.grid(axis="x", visible=False)
        _ax.spines[["top", "right", "left"]].set_visible(False)
        _ax.spines["bottom"].set_color("#d9d9d9")
        _ax.tick_params(axis="both", colors="#555555", length=0)
        _ax.legend(
            loc="lower center",
            bbox_to_anchor=(0.5, -0.22),
            ncol=4,
            frameon=False,
            fontsize=10,
        )
        _fig.tight_layout()
        return _fig

    def _latest_change(_segment, _provider):
        _provider_data = (
            market_share_abonnement_segment.filter(
                (pl.col("ms") == _segment) & (pl.col("tilbyder") == _provider)
            )
            .sort("ar")
        )
        _periods = _provider_data["ar"].to_list()
        _values = _provider_data["markedsandel"].to_list()
        if len(_values) < 2:
            return None
        return {
            "previous": _values[-2],
            "current": _values[-1],
            "delta": _values[-1] - _values[-2],
        }

    def _direction(_change):
        if _change["delta"] > 0.05:
            return "øker"
        if _change["delta"] < -0.05:
            return "går ned"
        return "er stabil"

    def _segment_text(_provider):
        _private = _latest_change("Privat", _provider)
        _business = _latest_change("Bedrift", _provider)
        if _private is None or _business is None:
            return ""
        return (
            f"<strong>{_provider}</strong> {_direction(_private)} i privatmarkedet "
            f"(<strong>{_percent(_private['previous'])}→{_percent(_private['current'])}</strong>) "
            f"og {_direction(_business)} i bedriftsmarkedet "
            f"(<strong>{_percent(_business['previous'])}→{_percent(_business['current'])}</strong>)."
        )

    def _last_two_periods():
        _years = sorted(market_share_abonnement_segment["ar"].unique().to_list())
        if len(_years) < 2:
            return None, None
        return _years[-2], _years[-1]

    _prev_period, _curr_period = _last_two_periods()
    _period_text = (
        f"Endring fra {_prev_period} til {_curr_period}"
        if _prev_period is not None and _curr_period is not None
        else "Endring mellom siste perioder"
    )
    fig3_private_fig = _plot_segment("Privat", 50)
    fig3_business_fig = _plot_segment("Bedrift")
    _private_export = export_menu_placeholder("fig3-privat-export")
    _business_export = export_menu_placeholder("fig3-bedrift-export")
    _summary = mo.Html(
        f"""
        <div style="
            margin: 8px auto 0;
            max-width: 820px;
            padding: 18px 26px;
            border: 1px solid #0b2b66;
            border-radius: 28px;
            background: #dcefd2;
            color: #000;
            font-size: 1.05rem;
            line-height: 1.45;
            text-align: center;
        ">
            <div style="margin: 0 0 12px; font-size: 1.2rem; font-weight: 700;">{_period_text}</div>
            {_segment_text("Lyse Tele (Ice)")}<br>
            {_segment_text("Telenor")}<br>
            {_segment_text("Telia")}<br>
            {_segment_text("Øvrige")}
        </div>
        """
    )

    mo.vstack(
        [
            mo.hstack(
                [
                    mo.vstack(
                        [
                            mo.Html(
                                '<div class="figure-heading-title">Privat</div>'
                            ),
                            figure_with_export(fig3_private_fig, _private_export),
                        ],
                        gap=0.5,
                    ),
                    mo.vstack(
                        [
                            mo.Html(
                                '<div class="figure-heading-title">Bedrift</div>'
                            ),
                            figure_with_export(fig3_business_fig, _business_export),
                        ],
                        gap=0.5,
                    ),
                ],
                justify="center",
                gap=2,
            ),
            _summary,
        ],
        gap=2,
    )
    return fig3_business_fig, fig3_private_fig


@app.cell
def _(
    export_menu,
    fig1_abonnement_fig,
    fig1_omsetning_fig,
    fig2_abonnement_fig,
    fig2_omsetning_fig,
    fig3_business_fig,
    fig3_private_fig,
    market_share_abonnement,
    market_share_abonnement_projection,
    market_share_abonnement_segment,
    market_share_omsetning,
    market_share_omsetning_projection,
    mo,
    pl,
    projection_export_data,
):
    _updates = [
        export_menu(
            "fig1-abonnement-export",
            market_share_abonnement,
            "figur-1-abonnement.xlsx",
            "Figur 1 - Markedsandeler basert på abonnement",
            "Abonnement",
            fig1_abonnement_fig,
            "figur-1-abonnement.png",
        ),
        export_menu(
            "fig1-omsetning-export",
            market_share_omsetning,
            "figur-1-omsetning.xlsx",
            "Figur 1 - Markedsandeler basert på omsetning",
            "Omsetning",
            fig1_omsetning_fig,
            "figur-1-omsetning.png",
        ),
        export_menu(
            "fig2-abonnement-export",
            projection_export_data(
                market_share_abonnement,
                market_share_abonnement_projection,
            ),
            "figur-2-abonnement-trend.xlsx",
            "Figur 2 - Lineær trend basert på abonnement",
            "Abonnement trend",
            fig2_abonnement_fig,
            "figur-2-abonnement-trend.png",
        ),
        export_menu(
            "fig2-omsetning-export",
            projection_export_data(
                market_share_omsetning,
                market_share_omsetning_projection,
            ),
            "figur-2-omsetning-trend.xlsx",
            "Figur 2 - Lineær trend basert på omsetning",
            "Omsetning trend",
            fig2_omsetning_fig,
            "figur-2-omsetning-trend.png",
        ),
        export_menu(
            "fig3-privat-export",
            market_share_abonnement_segment.filter(pl.col("ms") == "Privat"),
            "figur-3-privat.xlsx",
            "Figur 3 - Abonnement i privatmarkedet",
            "Privat",
            fig3_private_fig,
            "figur-3-privat.png",
        ),
        export_menu(
            "fig3-bedrift-export",
            market_share_abonnement_segment.filter(pl.col("ms") == "Bedrift"),
            "figur-3-bedrift.xlsx",
            "Figur 3 - Abonnement i bedriftsmarkedet",
            "Bedrift",
            fig3_business_fig,
            "figur-3-bedrift.png",
        ),
    ]
    mo.Html("\n".join(_updates))
    return


if __name__ == "__main__":
    app.run()
