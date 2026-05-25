from pathlib import Path

import matplotlib.pyplot as plt
import polars as pl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "dist" / "exports"
DATA = ROOT / "data" / "mobil.parquet"

ORDER = ["Telenor", "Telia", "Lyse Tele (Ice)", "Øvrige"]
COLORS = {
    "Telenor": "#156082",
    "Telia": "#7030a0",
    "Lyse Tele (Ice)": "#ffc000",
    "Øvrige": "#00b050",
}


def provider_group() -> pl.Expr:
    name = pl.col("fusnavn").str.to_lowercase()
    return (
        pl.when(name.str.contains("telenor"))
        .then(pl.lit("Telenor"))
        .when(name.str.contains("telia"))
        .then(pl.lit("Telia"))
        .when(
            name.is_in(["ice communication norge", "lyse tele"])
            | name.str.contains("lyse")
            | name.str.contains("ice")
        )
        .then(pl.lit("Lyse Tele (Ice)"))
        .otherwise(pl.lit("Øvrige"))
    )


def pct(value: float) -> str:
    return f"{value:.1f} %".replace(".", ",")


def market_share_abonnement(df: pl.LazyFrame) -> pl.DataFrame:
    return (
        df.filter(
            (pl.col("dk") == "Mobiltelefoni")
            & (pl.col("hg") == "Abonnement")
            & (pl.col("n1").is_in(["Fakturert", "Kontantkort"]))
            & (pl.col("n2") == "Ingen")
            & (pl.col("tp") == "Sum")
            & (pl.col("sk") == "Sluttbruker")
            & (pl.col("delar") == "Helår")
        )
        .with_columns(tilbyder=provider_group())
        .group_by("ar", "tilbyder")
        .agg(pl.col("svar").sum().alias("abonnement"))
        .with_columns(
            markedsandel=pl.col("abonnement")
            * 100
            / pl.col("abonnement").sum().over("ar")
        )
        .select("ar", "tilbyder", "markedsandel")
        .sort("ar", "tilbyder")
        .collect()
    )


def market_share_omsetning(df: pl.LazyFrame) -> pl.DataFrame:
    return (
        df.filter(
            (pl.col("dk") == "Mobiltelefoni")
            & (pl.col("hg") == "Inntekter")
            & (pl.col("tp") == "Sum")
            & (pl.col("sk") == "Sluttbruker")
            & (pl.col("delar") == "Helår")
        )
        .with_columns(tilbyder=provider_group())
        .group_by("ar", "tilbyder")
        .agg(pl.col("svar").sum().alias("omsetning"))
        .with_columns(
            markedsandel=pl.col("omsetning")
            * 100
            / pl.col("omsetning").sum().over("ar")
        )
        .select("ar", "tilbyder", "markedsandel")
        .sort("ar", "tilbyder")
        .collect()
    )


def market_share_segment(df: pl.LazyFrame) -> pl.DataFrame:
    return (
        df.filter(
            (pl.col("dk") == "Mobiltelefoni")
            & (pl.col("hg") == "Abonnement")
            & (pl.col("ms").is_in(["Privat", "Bedrift"]))
            & (pl.col("n1").is_in(["Fakturert", "Kontantkort"]))
            & (pl.col("n2") == "Ingen")
            & (pl.col("tp") == "Sum")
            & (pl.col("sk") == "Sluttbruker")
            & (pl.col("delar") == "Helår")
        )
        .with_columns(tilbyder=provider_group())
        .group_by("ar", "ms", "tilbyder")
        .agg(pl.col("svar").sum().alias("abonnement"))
        .with_columns(
            markedsandel=pl.col("abonnement")
            * 100
            / pl.col("abonnement").sum().over(["ar", "ms"])
        )
        .select("ar", "ms", "tilbyder", "markedsandel")
        .sort("ms", "ar", "tilbyder")
        .collect()
    )


def linear_projection(data: pl.DataFrame, periods_ahead: int = 3) -> pl.DataFrame:
    rows = []
    for provider in ORDER:
        provider_data = data.filter(pl.col("tilbyder") == provider).sort("ar")
        if provider_data.height < 2:
            continue

        xs = [float(x) for x in provider_data["ar"].to_list()]
        ys = [float(y) for y in provider_data["markedsandel"].to_list()]
        x_mean = sum(xs) / len(xs)
        y_mean = sum(ys) / len(ys)
        denominator = sum((x - x_mean) ** 2 for x in xs)
        slope = (
            sum((x - x_mean) * (y - y_mean) for x, y in zip(xs, ys)) / denominator
            if denominator
            else 0
        )
        last_year = int(max(xs))
        last_value = ys[-1]
        for year in range(last_year, last_year + periods_ahead + 1):
            rows.append(
                {
                    "ar": year,
                    "tilbyder": provider,
                    "markedsandel": last_value + slope * (year - last_year),
                }
            )
    return pl.DataFrame(rows)


def projection_export(actual: pl.DataFrame, projection: pl.DataFrame) -> pl.DataFrame:
    return (
        pl.concat(
            [
                actual.with_columns(
                    pl.lit("Historikk").alias("serie"),
                    pl.col("ar").cast(pl.Int64),
                ),
                projection.with_columns(
                    pl.lit("Lineær trend").alias("serie"),
                    pl.col("ar").cast(pl.Int64),
                ),
            ]
        )
        .select("serie", "ar", "tilbyder", "markedsandel")
        .sort("serie", "ar", "tilbyder")
    )


def plot_share(data: pl.DataFrame, upper_y: int, path: Path) -> None:
    fig, ax = plt.subplots(figsize=(7.2, 4.8), dpi=120)
    for provider in ORDER:
        provider_data = data.filter(pl.col("tilbyder") == provider).sort("ar")
        if provider_data.is_empty():
            continue
        xs = provider_data["ar"].to_list()
        ys = provider_data["markedsandel"].to_list()
        ax.plot(xs, ys, color=COLORS[provider], linewidth=3, solid_capstyle="round", label=provider)
        if len(ys) >= 2:
            ax.annotate(pct(ys[-2]), xy=(xs[-2], ys[-2]), xytext=(10, 6), textcoords="offset points", color="#404040", fontsize=9)
        ax.annotate(pct(ys[-1]), xy=(xs[-1], ys[-1]), xytext=(10, 6), textcoords="offset points", color="#404040", fontsize=9)
    style_axes(ax, data["ar"].unique().to_list(), upper_y, legend_y=-0.18)
    fig.savefig(path, format="png", dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def plot_projection(actual: pl.DataFrame, projection: pl.DataFrame, path: Path) -> None:
    fig, ax = plt.subplots(figsize=(7.8, 4.8), dpi=120)
    last_actual_year = int(actual["ar"].max())
    for provider in ORDER:
        actual_provider = actual.filter(pl.col("tilbyder") == provider).sort("ar")
        projection_provider = projection.filter(pl.col("tilbyder") == provider).sort("ar")
        if actual_provider.is_empty() or projection_provider.is_empty():
            continue
        actual_xs = actual_provider["ar"].to_list()
        actual_ys = actual_provider["markedsandel"].to_list()
        projection_xs = projection_provider["ar"].to_list()
        projection_ys = projection_provider["markedsandel"].to_list()
        ax.plot(actual_xs, actual_ys, color=COLORS[provider], linewidth=2.5, solid_capstyle="round", label=provider)
        ax.plot(projection_xs, projection_ys, color=COLORS[provider], linewidth=1.8, linestyle=":")
        latest = actual_provider.filter(pl.col("ar") == last_actual_year)
        if not latest.is_empty():
            ax.annotate(pct(latest["markedsandel"][0]), xy=(last_actual_year, latest["markedsandel"][0]), xytext=(8, 7), textcoords="offset points", color="#404040", fontsize=8)
        ax.annotate(pct(projection_ys[-1]), xy=(projection_xs[-1], projection_ys[-1]), xytext=(8, 0), textcoords="offset points", color="#202020", fontsize=8, bbox={"boxstyle": "square,pad=0.35", "facecolor": "white", "edgecolor": COLORS[provider], "linewidth": 0.8})
    years = sorted(set(actual["ar"].unique().to_list() + projection["ar"].unique().to_list()))
    style_axes(ax, years, 60, legend_y=-0.26, legend_fontsize=8)
    fig.savefig(path, format="png", dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def style_axes(ax, years, upper_y: int, legend_y: float, legend_fontsize: int = 10) -> None:
    ax.set_ylim(0, upper_y)
    ax.set_yticks(range(0, upper_y + 1, 10))
    ax.set_yticklabels([pct(value) for value in range(0, upper_y + 1, 10)])
    ax.set_xticks(sorted(years))
    ax.grid(axis="y", color="#d9d9d9", linewidth=0.8)
    ax.grid(axis="x", visible=False)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.spines["bottom"].set_color("#d9d9d9")
    ax.tick_params(axis="both", colors="#555555", length=0)
    ax.legend(loc="lower center", bbox_to_anchor=(0.5, legend_y), ncol=4, frameon=False, fontsize=legend_fontsize)
    ax.figure.tight_layout()


def write_excel(data: pl.DataFrame, path: Path, title: str, sheet_name: str) -> None:
    is_trend = "serie" in data.columns
    if is_trend:
        index_columns = [column for column in ["ms", "tilbyder"] if column in data.columns]
        history_years = sorted(data.filter(pl.col("serie") == "Historikk")["ar"].unique().to_list())
        last_history_year = max(history_years)
        trend_years = sorted(data.filter((pl.col("serie") == "Lineær trend") & (pl.col("ar") > last_history_year))["ar"].unique().to_list())
        export_data = (
            data.filter((pl.col("serie") == "Historikk") | ((pl.col("serie") == "Lineær trend") & (pl.col("ar") > last_history_year)))
            .with_columns((pl.col("markedsandel") / 100).alias("markedsandel"), pl.col("ar").cast(pl.Utf8))
            .pivot(values="markedsandel", index=index_columns, on="ar", aggregate_function="first")
            .select(index_columns + [str(year) for year in history_years + trend_years])
            .sort(index_columns)
        )
        years = history_years + trend_years
    else:
        index_columns = [column for column in ["ms", "tilbyder"] if column in data.columns]
        years = sorted(data["ar"].unique().to_list())
        export_data = (
            data.with_columns((pl.col("markedsandel") / 100).alias("markedsandel"), pl.col("ar").cast(pl.Utf8))
            .pivot(values="markedsandel", index=index_columns, on="ar", aggregate_function="first")
            .select(index_columns + [str(year) for year in years])
            .sort(index_columns)
        )
    headers = export_data.columns
    header_row = 4 if is_trend else 3
    first_data_row = header_row + 1
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color="0B2B66")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers)))
    ws.append([])
    if is_trend:
        ws.append([""] * len(headers))
        history_start = len(index_columns) + 1
        history_end = history_start + len(history_years) - 1
        trend_start = history_end + 1
        trend_end = trend_start + len(trend_years) - 1
        if history_years:
            ws.merge_cells(start_row=3, start_column=history_start, end_row=3, end_column=history_end)
            ws.cell(row=3, column=history_start).value = "Historikk"
        if trend_years:
            ws.merge_cells(start_row=3, start_column=trend_start, end_row=3, end_column=trend_end)
            ws.cell(row=3, column=trend_start).value = "Lineær trend"
    ws.append(headers)
    for row in export_data.iter_rows():
        ws.append(list(row))
    header_fill = PatternFill("solid", fgColor="0B2B66")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9D9D9")
    medium_blue = Side(style="medium", color="0B2B66")
    border = Border(bottom=thin)
    if is_trend:
        history_fill = PatternFill("solid", fgColor="E9F2DF")
        trend_fill = PatternFill("solid", fgColor="DDEBF7")
        for cell in ws[3]:
            cell.fill = history_fill
            cell.font = Font(bold=True, color="000000")
            cell.alignment = Alignment(horizontal="center")
            cell.border = Border(top=medium_blue, bottom=thin)
        for column_index in range(trend_start, trend_end + 1):
            cell = ws.cell(row=3, column=column_index)
            cell.fill = trend_fill
            cell.border = Border(top=medium_blue, bottom=thin)
    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        if is_trend and cell.column >= trend_start:
            cell.fill = PatternFill("solid", fgColor="1F4E79")
    for row in ws.iter_rows(min_row=first_data_row, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = border
            if cell.column_letter == "A":
                cell.alignment = Alignment(horizontal="left")
            if headers[cell.column - 1] in [str(year) for year in years]:
                cell.number_format = "0.0%"
    table_ref = f"A{header_row}:{ws.cell(ws.max_row, ws.max_column).coordinate}"
    for row_index in range(first_data_row, ws.max_row + 1):
        if (row_index - first_data_row) % 2 == 0:
            for cell in ws[row_index]:
                cell.fill = PatternFill("solid", fgColor="EEF3F8")
        if is_trend:
            for column_index in range(len(index_columns) + 1, history_end + 1):
                ws.cell(row=row_index, column=column_index).fill = PatternFill("solid", fgColor="F3F8EC")
            for column_index in range(trend_start, trend_end + 1):
                ws.cell(row=row_index, column=column_index).fill = PatternFill("solid", fgColor="EAF3FB")
            left_cell = ws.cell(row=row_index, column=trend_start)
            left_cell.border = Border(left=medium_blue, right=left_cell.border.right, top=left_cell.border.top, bottom=left_cell.border.bottom)
    ws.freeze_panes = f"A{first_data_row}"
    ws.auto_filter.ref = table_ref
    for column_index, header in enumerate(headers, start=1):
        max_length = len(str(header))
        for row_index in range(first_data_row, ws.max_row + 1):
            value = ws.cell(row=row_index, column=column_index).value
            max_length = max(max_length, len(str(value)) if value is not None else 0)
        ws.column_dimensions[ws.cell(row=header_row, column=column_index).column_letter].width = min(max(max_length + 2, 12), 28)
    wb.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    df = pl.scan_parquet(DATA)
    ab = market_share_abonnement(df)
    om = market_share_omsetning(df)
    seg = market_share_segment(df)
    ab_proj = linear_projection(ab)
    om_proj = linear_projection(om)
    plot_share(ab, 60, OUT / "figur-1-abonnement.png")
    plot_share(om, 60, OUT / "figur-1-omsetning.png")
    plot_projection(ab, ab_proj, OUT / "figur-2-abonnement-trend.png")
    plot_projection(om, om_proj, OUT / "figur-2-omsetning-trend.png")
    plot_share(seg.filter(pl.col("ms") == "Privat"), 50, OUT / "figur-3-privat.png")
    plot_share(seg.filter(pl.col("ms") == "Bedrift"), 60, OUT / "figur-3-bedrift.png")
    write_excel(ab, OUT / "figur-1-abonnement.xlsx", "Figur 1 - Markedsandeler basert på abonnement", "Abonnement")
    write_excel(om, OUT / "figur-1-omsetning.xlsx", "Figur 1 - Markedsandeler basert på omsetning", "Omsetning")
    write_excel(projection_export(ab, ab_proj), OUT / "figur-2-abonnement-trend.xlsx", "Figur 2 - Lineær trend basert på abonnement", "Abonnement trend")
    write_excel(projection_export(om, om_proj), OUT / "figur-2-omsetning-trend.xlsx", "Figur 2 - Lineær trend basert på omsetning", "Omsetning trend")
    write_excel(seg.filter(pl.col("ms") == "Privat"), OUT / "figur-3-privat.xlsx", "Figur 3 - Abonnement i privatmarkedet", "Privat")
    write_excel(seg.filter(pl.col("ms") == "Bedrift"), OUT / "figur-3-bedrift.xlsx", "Figur 3 - Abonnement i bedriftsmarkedet", "Bedrift")


if __name__ == "__main__":
    main()
