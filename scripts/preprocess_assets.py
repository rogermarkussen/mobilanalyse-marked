from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import polars as pl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "mobil.parquet"
FIGURE_DIR = ROOT / "dist" / "assets" / "figures"
EXPORT_DIR = ROOT / "dist" / "assets" / "exports"
GENERATED = ROOT / "mobilanalyse_generated.py"

ORDER = ["Telenor", "Telia", "Lyse Tele (Ice)", "Øvrige"]
SUMMARY_ORDER = ["Lyse Tele (Ice)", "Telenor", "Telia", "Øvrige"]
COLORS = {
    "Telenor": "#156082",
    "Telia": "#7030a0",
    "Lyse Tele (Ice)": "#ffc000",
    "Øvrige": "#00b050",
}


def provider_group() -> pl.Expr:
    name = pl.col("fusnavn").str.to_lowercase()
    return (
        pl.when(name.str.contains("telenor"))
        .then(pl.lit("Telenor"))
        .when(name.str.contains("telia"))
        .then(pl.lit("Telia"))
        .when(
            name.is_in(["ice communication norge", "lyse tele"])
            | name.str.contains("lyse")
            | name.str.contains("ice")
        )
        .then(pl.lit("Lyse Tele (Ice)"))
        .otherwise(pl.lit("Øvrige"))
    )


def pct(value: float) -> str:
    return f"{value:.1f} %".replace(".", ",")


def market_share_abonnement(df: pl.LazyFrame) -> pl.DataFrame:
    return (
        df.filter(
            (pl.col("dk") == "Mobiltelefoni")
            & (pl.col("hg") == "Abonnement")
            & (pl.col("n1").is_in(["Fakturert", "Kontantkort"]))
            & (pl.col("n2") == "Ingen")
            & (pl.col("tp") == "Sum")
            & (pl.col("sk") == "Sluttbruker")
            & (pl.col("delar") == "Helår")
        )
        .with_columns(tilbyder=provider_group())
        .group_by("ar", "tilbyder")
        .agg(pl.col("svar").sum().alias("abonnement"))
        .with_columns(
            markedsandel=pl.col("abonnement")
            * 100
            / pl.col("abonnement").sum().over("ar")
        )
        .select("ar", "tilbyder", "markedsandel")
        .sort("ar", "tilbyder")
        .collect()
    )


def market_share_omsetning(df: pl.LazyFrame) -> pl.DataFrame:
    return (
        df.filter(
            (
                ((pl.col("dk") == "Mobiltelefoni") & (pl.col("sk") == "Sluttbruker"))
                | (pl.col("dk") == "Roaming")
            )
            & (pl.col("hg") == "Inntekter")
            & (pl.col("tp") == "Sum")
            & (pl.col("delar") == "Helår")
        )
        .with_columns(tilbyder=provider_group())
        .group_by("ar", "tilbyder")
        .agg(pl.col("svar").sum().alias("omsetning"))
        .with_columns(
            markedsandel=pl.col("omsetning")
            * 100
            / pl.col("omsetning").sum().over("ar")
        )
        .select("ar", "tilbyder", "markedsandel")
        .sort("ar", "tilbyder")
        .collect()
    )


def market_share_segment(df: pl.LazyFrame) -> pl.DataFrame:
    return (
        df.filter(
            (pl.col("dk") == "Mobiltelefoni")
            & (pl.col("hg") == "Abonnement")
            & (pl.col("ms").is_in(["Privat", "Bedrift"]))
            & (pl.col("n1").is_in(["Fakturert", "Kontantkort"]))
            & (pl.col("n2") == "Ingen")
            & (pl.col("tp") == "Sum")
            & (pl.col("sk") == "Sluttbruker")
            & (pl.col("delar") == "Helår")
        )
        .with_columns(tilbyder=provider_group())
        .group_by("ar", "ms", "tilbyder")
        .agg(pl.col("svar").sum().alias("abonnement"))
        .with_columns(
            markedsandel=pl.col("abonnement")
            * 100
            / pl.col("abonnement").sum().over(["ar", "ms"])
        )
        .select("ar", "ms", "tilbyder", "markedsandel")
        .sort("ms", "ar", "tilbyder")
        .collect()
    )


def linear_projection(data: pl.DataFrame, periods_ahead: int = 3) -> pl.DataFrame:
    rows = []
    for provider in ORDER:
        provider_data = data.filter(pl.col("tilbyder") == provider).sort("ar")
        if provider_data.height < 2:
            continue

        slope = linear_slope(provider_data)
        years = provider_data["ar"].to_list()
        values = provider_data["markedsandel"].to_list()
        last_year = int(years[-1])
        last_value = float(values[-1])
        for year in range(last_year, last_year + periods_ahead + 1):
            rows.append(
                {
                    "ar": year,
                    "tilbyder": provider,
                    "markedsandel": last_value + slope * (year - last_year),
                }
            )
    return pl.DataFrame(rows)


def linear_slope(provider_data: pl.DataFrame) -> float:
    xs = [float(x) for x in provider_data["ar"].to_list()]
    ys = [float(y) for y in provider_data["markedsandel"].to_list()]
    x_mean = sum(xs) / len(xs)
    y_mean = sum(ys) / len(ys)
    denominator = sum((x - x_mean) ** 2 for x in xs)
    if denominator == 0:
        return 0.0
    return sum((x - x_mean) * (y - y_mean) for x, y in zip(xs, ys)) / denominator


def style_axes(ax, years: list[int], upper_y: int, legend_y: float, legend_fontsize: int = 10) -> None:
    ax.set_ylim(0, upper_y)
    ax.set_yticks(range(0, upper_y + 1, 10))
    ax.set_yticklabels([pct(value) for value in range(0, upper_y + 1, 10)])
    ax.set_xticks(sorted(years))
    ax.grid(axis="y", color="#d9d9d9", linewidth=0.8)
    ax.grid(axis="x", visible=False)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.spines["bottom"].set_color("#d9d9d9")
    ax.tick_params(axis="both", colors="#555555", length=0)
    ax.legend(
        loc="lower center",
        bbox_to_anchor=(0.5, legend_y),
        ncol=4,
        frameon=False,
        fontsize=legend_fontsize,
    )
    ax.figure.tight_layout()


def plot_share(data: pl.DataFrame, upper_y: int, path: Path) -> None:
    fig, ax = plt.subplots(figsize=(7.2, 4.8), dpi=120)
    for provider in ORDER:
        provider_data = data.filter(pl.col("tilbyder") == provider).sort("ar")
        if provider_data.is_empty():
            continue
        xs = provider_data["ar"].to_list()
        ys = provider_data["markedsandel"].to_list()
        ax.plot(
            xs,
            ys,
            color=COLORS[provider],
            linewidth=3,
            solid_capstyle="round",
            label=provider,
        )
        if len(ys) >= 2:
            ax.annotate(
                pct(ys[-2]),
                xy=(xs[-2], ys[-2]),
                xytext=(10, 6),
                textcoords="offset points",
                color="#404040",
                fontsize=9,
            )
        ax.annotate(
            pct(ys[-1]),
            xy=(xs[-1], ys[-1]),
            xytext=(10, 6),
            textcoords="offset points",
            color="#404040",
            fontsize=9,
        )
    style_axes(ax, data["ar"].unique().to_list(), upper_y, legend_y=-0.18)
    fig.savefig(path, format="png", dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def plot_projection(actual: pl.DataFrame, projection: pl.DataFrame, path: Path) -> None:
    fig, ax = plt.subplots(figsize=(7.8, 4.8), dpi=120)
    last_actual_year = int(actual["ar"].max())
    for provider in ORDER:
        actual_provider = actual.filter(pl.col("tilbyder") == provider).sort("ar")
        projection_provider = projection.filter(pl.col("tilbyder") == provider).sort("ar")
        if actual_provider.is_empty() or projection_provider.is_empty():
            continue
        actual_xs = actual_provider["ar"].to_list()
        actual_ys = actual_provider["markedsandel"].to_list()
        projection_xs = projection_provider["ar"].to_list()
        projection_ys = projection_provider["markedsandel"].to_list()
        ax.plot(
            actual_xs,
            actual_ys,
            color=COLORS[provider],
            linewidth=2.5,
            solid_capstyle="round",
            label=provider,
        )
        ax.plot(
            projection_xs,
            projection_ys,
            color=COLORS[provider],
            linewidth=1.8,
            linestyle=":",
        )

        latest = actual_provider.filter(pl.col("ar") == last_actual_year)
        if not latest.is_empty():
            ax.annotate(
                pct(latest["markedsandel"][0]),
                xy=(last_actual_year, latest["markedsandel"][0]),
                xytext=(8, 7),
                textcoords="offset points",
                color="#404040",
                fontsize=8,
            )
        ax.annotate(
            pct(projection_ys[-1]),
            xy=(projection_xs[-1], projection_ys[-1]),
            xytext=(8, 0),
            textcoords="offset points",
            color="#202020",
            fontsize=8,
            bbox={
                "boxstyle": "square,pad=0.35",
                "facecolor": "white",
                "edgecolor": COLORS[provider],
                "linewidth": 0.8,
            },
        )
    years = sorted(set(actual["ar"].unique().to_list() + projection["ar"].unique().to_list()))
    style_axes(ax, years, 60, legend_y=-0.22, legend_fontsize=9)
    fig.savefig(path, format="png", dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def latest_change(data: pl.DataFrame, provider: str) -> dict[str, float | int] | None:
    provider_data = data.filter(pl.col("tilbyder") == provider).sort("ar")
    years = provider_data["ar"].to_list()
    values = provider_data["markedsandel"].to_list()
    if len(values) < 2:
        return None
    return {
        "previous_period": int(years[-2]),
        "current_period": int(years[-1]),
        "previous": float(values[-2]),
        "current": float(values[-1]),
        "delta": float(values[-1] - values[-2]),
    }


def direction(change: dict[str, float | int], stable_text: str = "er om lag uendret") -> str:
    delta = float(change["delta"])
    if delta > 0.05:
        return "øker"
    if delta < -0.05:
        return "går ned"
    return stable_text


def pair_text(ab: pl.DataFrame, om: pl.DataFrame, provider: str) -> str:
    ab_change = latest_change(ab, provider)
    om_change = latest_change(om, provider)
    if ab_change is None or om_change is None:
        return ""
    return (
        f"<strong>{provider}</strong> {direction(om_change)} i omsetningsandel "
        f"(<strong>{pct(float(om_change['previous']))}&rarr;{pct(float(om_change['current']))}</strong>) "
        f"og {direction(ab_change)} i abonnementsandel "
        f"(<strong>{pct(float(ab_change['previous']))}&rarr;{pct(float(ab_change['current']))}</strong>)."
    )


def period_text(*frames: pl.DataFrame) -> str:
    years: set[int] = set()
    for frame in frames:
        years.update(int(year) for year in frame["ar"].to_list())
    sorted_years = sorted(years)
    if len(sorted_years) < 2:
        return "Endring mellom siste perioder"
    return f"Endring fra {sorted_years[-2]} til {sorted_years[-1]}"


def segment_change(data: pl.DataFrame, segment: str, provider: str) -> dict[str, float | int] | None:
    return latest_change(
        data.filter((pl.col("ms") == segment) & (pl.col("tilbyder") == provider)),
        provider,
    )


def segment_text(data: pl.DataFrame, provider: str) -> str:
    private = segment_change(data, "Privat", provider)
    business = segment_change(data, "Bedrift", provider)
    if private is None or business is None:
        return ""
    return (
        f"<strong>{provider}</strong> {direction(private, 'er stabil')} i privatmarkedet "
        f"(<strong>{pct(float(private['previous']))}&rarr;{pct(float(private['current']))}</strong>) "
        f"og {direction(business, 'er stabil')} i bedriftsmarkedet "
        f"(<strong>{pct(float(business['previous']))}&rarr;{pct(float(business['current']))}</strong>)."
    )


def threshold_text(data: pl.DataFrame, provider: str, threshold: float) -> str:
    provider_data = data.filter(pl.col("tilbyder") == provider).sort("ar")
    if provider_data.height < 2:
        return f"{provider} har ikke nok historikk til å beregne lineær trend."

    slope = linear_slope(provider_data)
    last_year = int(provider_data["ar"][-1])
    last_value = float(provider_data["markedsandel"][-1])
    if last_value >= threshold:
        return f"{provider} er allerede over {threshold:.0f} % omsetningsandel."
    if slope <= 0:
        return f"{provider} når ikke {threshold:.0f} % omsetningsandel med lineær trend."

    projected_year = last_year + (threshold - last_value) / slope
    return f"{provider} når {threshold:.0f} % av omsetningen rundt {int(projected_year + 0.999)}."


def export_matrix(data: pl.DataFrame) -> tuple[list[str], list[list[str | float | None]]]:
    years = sorted(int(year) for year in data["ar"].unique().to_list())
    pivoted = (
        data.with_columns(
            (pl.col("markedsandel") / 100).alias("markedsandel"),
            pl.col("ar").cast(pl.Utf8),
        )
        .pivot(
            values="markedsandel",
            index="tilbyder",
            on="ar",
            aggregate_function="first",
        )
        .select(["tilbyder"] + [str(year) for year in years])
        .sort("tilbyder")
    )
    return pivoted.columns, [list(row) for row in pivoted.iter_rows()]


def trend_export_matrix(
    actual: pl.DataFrame,
    projection: pl.DataFrame,
) -> tuple[list[str], list[int], list[int], list[list[str | float | None]]]:
    history_years = sorted(int(year) for year in actual["ar"].unique().to_list())
    last_history_year = max(history_years)
    trend_years = sorted(
        int(year)
        for year in projection.filter(pl.col("ar") > last_history_year)["ar"].unique().to_list()
    )

    history = actual.with_columns(
        (pl.col("markedsandel") / 100).alias("markedsandel"),
        pl.col("ar").cast(pl.Utf8),
    )
    trend = projection.filter(pl.col("ar") > last_history_year).with_columns(
        (pl.col("markedsandel") / 100).alias("markedsandel"),
        pl.col("ar").cast(pl.Utf8),
    )

    combined = pl.concat([history, trend], how="diagonal")
    columns = ["tilbyder"] + [str(year) for year in history_years + trend_years]
    pivoted = (
        combined.pivot(
            values="markedsandel",
            index="tilbyder",
            on="ar",
            aggregate_function="first",
        )
        .select(columns)
        .sort("tilbyder")
    )
    return columns, history_years, trend_years, [list(row) for row in pivoted.iter_rows()]


def style_workbook(
    wb: Workbook,
    title: str,
    columns: list[str],
    rows: list[list[str | float | None]],
    *,
    history_years: list[int] | None = None,
    trend_years: list[int] | None = None,
) -> None:
    ws = wb.active
    ws.title = "Markedsandeler"

    dark_blue = "0B2B66"
    history_fill = "E9F2DF"
    trend_fill = "DDEBF7"
    zebra_fill = "EEF3F8"
    border_gray = "D9D9D9"
    thin = Side(style="thin", color=border_gray)
    medium_blue = Side(style="medium", color=dark_blue)

    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color=dark_blue)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))

    header_row = 4 if history_years is not None and trend_years is not None else 3
    first_data_row = header_row + 1

    if history_years is not None and trend_years is not None:
        history_start = 2
        history_end = history_start + len(history_years) - 1
        trend_start = history_end + 1
        trend_end = trend_start + len(trend_years) - 1

        ws.merge_cells(start_row=3, start_column=history_start, end_row=3, end_column=history_end)
        ws.cell(row=3, column=history_start).value = "Historikk"
        ws.cell(row=3, column=history_start).fill = PatternFill("solid", fgColor=history_fill)
        ws.cell(row=3, column=history_start).font = Font(bold=True, color="000000")
        ws.cell(row=3, column=history_start).alignment = Alignment(horizontal="center")

        if trend_years:
            ws.merge_cells(start_row=3, start_column=trend_start, end_row=3, end_column=trend_end)
            ws.cell(row=3, column=trend_start).value = "Lineær trend"
            ws.cell(row=3, column=trend_start).fill = PatternFill("solid", fgColor=trend_fill)
            ws.cell(row=3, column=trend_start).font = Font(bold=True, color="000000")
            ws.cell(row=3, column=trend_start).alignment = Alignment(horizontal="center")

        for column_index in range(history_start, history_end + 1):
            ws.cell(row=3, column=column_index).fill = PatternFill("solid", fgColor=history_fill)
            ws.cell(row=3, column=column_index).border = Border(top=medium_blue, bottom=thin)
        for column_index in range(trend_start, trend_end + 1):
            ws.cell(row=3, column=column_index).fill = PatternFill("solid", fgColor=trend_fill)
            ws.cell(row=3, column=column_index).border = Border(top=medium_blue, bottom=thin)

    for column_index, column in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=column_index, value=column)
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(bottom=thin)

    for row_values in rows:
        ws.append(row_values)

    for row_index in range(first_data_row, ws.max_row + 1):
        if (row_index - first_data_row) % 2 == 0:
            for cell in ws[row_index]:
                cell.fill = PatternFill("solid", fgColor=zebra_fill)
        for column_index in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_index, column=column_index)
            cell.border = Border(bottom=thin)
            if column_index == 1:
                cell.alignment = Alignment(horizontal="left")
            else:
                cell.number_format = "0.0%"
                cell.alignment = Alignment(horizontal="right")

            if history_years is not None and trend_years is not None:
                if 2 <= column_index <= history_end:
                    cell.fill = PatternFill("solid", fgColor=history_fill)
                if trend_years and trend_start <= column_index <= trend_end:
                    cell.fill = PatternFill("solid", fgColor=trend_fill)
                if trend_years and column_index == trend_start:
                    cell.border = Border(
                        left=medium_blue,
                        bottom=thin,
                    )

    if history_years is not None and trend_years:
        for row_index in range(3, ws.max_row + 1):
            cell = ws.cell(row=row_index, column=trend_start)
            existing = cell.border
            cell.border = Border(
                left=medium_blue,
                right=existing.right,
                top=existing.top,
                bottom=existing.bottom,
            )

    ws.freeze_panes = f"A{first_data_row}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"

    for column_index, column in enumerate(columns, start=1):
        values = [column]
        for row_index in range(first_data_row, ws.max_row + 1):
            values.append(str(ws.cell(row=row_index, column=column_index).value or ""))
        width = max(len(str(value)) for value in values) + 2
        ws.column_dimensions[get_column_letter(column_index)].width = min(max(width, 12), 28)


def write_share_excel(data: pl.DataFrame, path: Path, title: str) -> None:
    columns, rows = export_matrix(data)
    wb = Workbook()
    style_workbook(wb, title, columns, rows)
    wb.save(path)


def write_trend_excel(actual: pl.DataFrame, projection: pl.DataFrame, path: Path, title: str) -> None:
    columns, history_years, trend_years, rows = trend_export_matrix(actual, projection)
    wb = Workbook()
    style_workbook(
        wb,
        title,
        columns,
        rows,
        history_years=history_years,
        trend_years=trend_years,
    )
    wb.save(path)


def write_generated(ab: pl.DataFrame, om: pl.DataFrame, seg: pl.DataFrame) -> None:
    figure_paths = {
        "fig1_abonnement": "assets/figures/figur-1-abonnement.png",
        "fig1_omsetning": "assets/figures/figur-1-omsetning.png",
        "fig2_abonnement": "assets/figures/figur-2-abonnement-trend.png",
        "fig2_omsetning": "assets/figures/figur-2-omsetning-trend.png",
        "fig3_privat": "assets/figures/figur-3-privat.png",
        "fig3_bedrift": "assets/figures/figur-3-bedrift.png",
    }
    export_paths = {
        "fig1_abonnement": "assets/exports/figur-1-abonnement.xlsx",
        "fig1_omsetning": "assets/exports/figur-1-omsetning.xlsx",
        "fig2_abonnement": "assets/exports/figur-2-abonnement-trend.xlsx",
        "fig2_omsetning": "assets/exports/figur-2-omsetning-trend.xlsx",
        "fig3_privat": "assets/exports/figur-3-privat.xlsx",
        "fig3_bedrift": "assets/exports/figur-3-bedrift.xlsx",
    }
    content = f'''# This file is generated by scripts/preprocess_assets.py.

FIGURE_PATHS = {figure_paths!r}
EXPORT_PATHS = {export_paths!r}

FIG1_PERIOD_TEXT = {period_text(ab, om)!r}
FIG1_SUMMARY_LINES = {[pair_text(ab, om, provider) for provider in SUMMARY_ORDER]!r}

FIG2_NOTE = {threshold_text(om, "Lyse Tele (Ice)", 20)!r}

FIG3_PERIOD_TEXT = {period_text(seg)!r}
FIG3_SUMMARY_LINES = {[segment_text(seg, provider) for provider in SUMMARY_ORDER]!r}
'''
    GENERATED.write_text(content, encoding="utf-8")


def main() -> None:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    df = pl.scan_parquet(DATA)

    ab = market_share_abonnement(df)
    om = market_share_omsetning(df)
    seg = market_share_segment(df)
    ab_projection = linear_projection(ab)
    om_projection = linear_projection(om)

    plot_share(ab, 60, FIGURE_DIR / "figur-1-abonnement.png")
    plot_share(om, 60, FIGURE_DIR / "figur-1-omsetning.png")
    plot_projection(ab, ab_projection, FIGURE_DIR / "figur-2-abonnement-trend.png")
    plot_projection(om, om_projection, FIGURE_DIR / "figur-2-omsetning-trend.png")
    plot_share(seg.filter(pl.col("ms") == "Privat"), 50, FIGURE_DIR / "figur-3-privat.png")
    plot_share(seg.filter(pl.col("ms") == "Bedrift"), 60, FIGURE_DIR / "figur-3-bedrift.png")

    write_share_excel(
        ab,
        EXPORT_DIR / "figur-1-abonnement.xlsx",
        "Figur 1 - Markedsandeler basert på abonnement",
    )
    write_share_excel(
        om,
        EXPORT_DIR / "figur-1-omsetning.xlsx",
        "Figur 1 - Markedsandeler basert på omsetning",
    )
    write_trend_excel(
        ab,
        ab_projection,
        EXPORT_DIR / "figur-2-abonnement-trend.xlsx",
        "Figur 2 - Lineær trend basert på abonnement",
    )
    write_trend_excel(
        om,
        om_projection,
        EXPORT_DIR / "figur-2-omsetning-trend.xlsx",
        "Figur 2 - Lineær trend basert på omsetning",
    )
    write_share_excel(
        seg.filter(pl.col("ms") == "Privat"),
        EXPORT_DIR / "figur-3-privat.xlsx",
        "Figur 3 - Abonnement i privatmarkedet",
    )
    write_share_excel(
        seg.filter(pl.col("ms") == "Bedrift"),
        EXPORT_DIR / "figur-3-bedrift.xlsx",
        "Figur 3 - Abonnement i bedriftsmarkedet",
    )

    write_generated(ab, om, seg)


if __name__ == "__main__":
    main()
